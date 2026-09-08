#!/usr/bin/env python
"""Извлечение ядра из книги xlsx с настоящими именами вкладок."""
import io, re
from collections import Counter, defaultdict
import openpyxl

SRC = r"C:\Users\Рома\Downloads\Недвижемость.xlsx"
wb = openpyxl.load_workbook(SRC, read_only=True, data_only=True)

def topic_lang(sheet):
    """Имя вкладки в xlsx обрезано до 31 символа, поэтому язык берём по началу:
    «Нерухомість…» — украинский, «Недвижимость…» — русский."""
    s = sheet.lower()
    lang = "uk" if s.startswith("нерухом") else "ru"
    for mark, name in (("іспан", "Испания"), ("испан", "Испания"),
                       ("кипр", "Северный Кипр"), ("кіпр", "Северный Кипр"),
                       ("бали", "Бали"), ("балі", "Бали"), ("моря", "У моря"),
                       ("алани", "Алания"), ("аланія", "Алания"),
                       ("турция", "Турция"), ("турція", "Турция")):
        if mark in s:
            return name, lang
    return "Прочее", lang

rows = []
for name in wb.sheetnames:
    if name == "Лендинги":
        continue
    topic, lang = topic_lang(name)
    ws = wb[name]
    data = list(ws.iter_rows(values_only=True))
    if not data:
        continue
    header = [str(c or "").strip() for c in data[0]]
    if not header or header[0] != "Ключевые фразы":
        continue
    idx = {h: i for i, h in enumerate(header)}
    def col(r, h):
        i = idx.get(h)
        return str(r[i]).strip() if i is not None and i < len(r) and r[i] is not None else ""
    for r in data[1:]:
        kw = str(r[0]).strip() if r and r[0] else ""
        if not kw:
            continue
        rows.append({"keyword": kw, "sheet": name, "topic": topic, "lang": lang,
                     "topo": col(r, "Топонимы"), "intent": col(r, "Интент"),
                     "difficulty": col(r, "Сложность"), "freq": col(r, "Частотность"),
                     "words": col(r, "Количество слов в ключевой фразе")})

seen = set(); uniq = []
for r in rows:
    k = (r["keyword"].lower(), r["lang"])
    if k not in seen:
        seen.add(k); uniq.append(r)

print(f"строк {len(rows)}, уникальных пар ключ+язык {len(uniq)}\n")
print(f"{'Тема':<16}{'RU':>6}{'UK':>6}{'Всего':>8}")
agg = defaultdict(lambda: Counter())
for r in uniq:
    agg[r["topic"]][r["lang"]] += 1
for topic, c in sorted(agg.items(), key=lambda kv: -sum(kv[1].values())):
    print(f"{topic:<16}{c['ru']:>6}{c['uk']:>6}{sum(c.values()):>8}")

with io.open("data/ads/old-core-clean.tsv", "w", encoding="utf-8") as fh:
    fh.write("keyword\tlang\ttopic\tsheet\ttopo\tintent\tdifficulty\tfreq_sheet\twords\n")
    for r in uniq:
        fh.write("\t".join([r["keyword"], r["lang"], r["topic"], r["sheet"], r["topo"],
                            r["intent"], r["difficulty"], r["freq"], r["words"]]) + "\n")
print("\nФайл: data/ads/old-core-clean.tsv")
